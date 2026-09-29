import os
import re
import sys

def evaluate_ats_compliance(file_path):
    print("=" * 70)
    print(f"SIH ATS COMPLIANCE AUDIT: {os.path.basename(file_path)}")
    print("=" * 70)
    
    if not os.path.exists(file_path):
        print(f"Error: File not found at {file_path}")
        return
        
    ext = os.path.splitext(file_path)[1].lower()
    text_content = ""
    page_count = 0
    
    if ext == ".pptx":
        try:
            import pptx
            prs = pptx.Presentation(file_path)
            page_count = len(prs.slides)
            for s_idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            if p.text.strip():
                                slide_text.append(p.text.strip())
                text_content += f"\n--- SLIDE {s_idx+1} ---\n" + "\n".join(slide_text)
        except Exception as e:
            print(f"Error reading PPTX: {e}")
            return
            
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            page_count = len(reader.pages)
            for p_idx, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                text_content += f"\n--- PAGE {p_idx+1} ---\n" + extracted
        except Exception as e:
            print(f"Error reading PDF: {e}")
            return
    else:
        print(f"Unsupported file format: {ext}. Must be .pptx or .pdf")
        return

    word_count = len(text_content.split())
    print(f"[*] Document Format: {ext.upper()}")
    print(f"[*] Total Slides/Pages: {page_count}")
    print(f"[*] Extracted Word Count: {word_count} words")
    
    # 1. Page Count Score (Max 15 pts)
    page_score = 0
    if page_count == 6:
        page_score = 15
        print("  [PASS] Page Count: Exactly 6 slides (15/15 pts)")
    elif page_count == 7:
        page_score = 8
        print(f"  [WARN] Page Count: {page_count} slides. SIH strictly requires 6 slides! (8/15 pts)")
    else:
        page_score = 3
        print(f"  [FAIL] Page Count: {page_count} slides. Strict 6-slide violation! (3/15 pts)")

    # 2. Machine-Readable Text Check (Max 15 pts)
    text_score = 0
    if word_count > 600:
        text_score = 15
        print(f"  [PASS] Machine-Readable Text: Rich, searchable text layer ({word_count} words) (15/15 pts)")
    elif word_count > 250:
        text_score = 10
        print(f"  [PASS] Machine-Readable Text: Acceptable word count ({word_count} words) (10/15 pts)")
    elif word_count > 50:
        text_score = 5
        print(f"  [WARN] Machine-Readable Text: Sparse text ({word_count} words). High risk of ATS drop! (5/15 pts)")
    else:
        text_score = 0
        print("  [FAIL] Machine-Readable Text: Less than 50 words! Document is mostly raster images. ATS will fail! (0/15 pts)")

    # 3. Official SIH Heading Rubric Check (Max 30 pts - 5 pts each)
    rubric_sections = [
        ("Slide 1 Metadata", [r"SIH26108", r"ZenicX", r"SIH-10", r"Smart India Hackathon", r"SOFTWARE"]),
        ("Slide 2 Proposed Solution", [r"IDEA TITLE", r"PROPOSED SOLUTION", r"INNOVATION"]),
        ("Slide 3 Technical Approach", [r"TECHNICAL APPROACH", r"METHODOLOGY|PROCESS|ARCHITECTURE"]),
        ("Slide 4 Feasibility & Viability", [r"FEASIBILITY", r"VIABILITY|CHALLENGES|RISKS|MITIGATION"]),
        ("Slide 5 Impact & Benefits", [r"IMPACT", r"BENEFITS|ECONOMIC|SOCIAL"]),
        ("Slide 6 Research & References", [r"RESEARCH", r"REFERENCES|CITATIONS"])
    ]
    
    heading_score = 0
    print("\n--- OFFICIAL SIH HEADING COMPLIANCE ---")
    for sec_name, patterns in rubric_sections:
        matches = [p for p in patterns if re.search(p, text_content, re.IGNORECASE)]
        if len(matches) == len(patterns):
            heading_score += 5
            print(f"  [PASS] {sec_name}: 100% matched ({5}/5 pts)")
        elif len(matches) > 0:
            heading_score += 3
            print(f"  [PARTIAL] {sec_name}: Matched {matches} ({3}/5 pts)")
        else:
            print(f"  [FAIL] {sec_name}: Missing required headings! (0/5 pts)")

    # 4. Domain & Technical Keyword Density (Max 40 pts)
    keyword_groups = {
        "Statutory & BIS Domain (10 pts)": [
            "Bureau of Indian Standards", "BIS", "Indian Standards", "IS Code",
            "Quality Control Orders", "QCO", "BIS Act 2016", "Section 16", "ISI mark"
        ],
        "Information Retrieval & Architecture (10 pts)": [
            "Hybrid Retrieval", "BM25", "Dense", "Embeddings", "Reciprocal Rank Fusion",
            "RRF", "Cross-Encoder", "Rerank", "ONNX", "FastAPI", "Next.js"
        ],
        "Zero-Hallucination & Performance (10 pts)": [
            "Hallucination", "Deterministic", "Clause", "560ms", "Latency",
            "22,446", "CPU Inference", "Vector DB"
        ],
        "National Digital India Infrastructure (10 pts)": [
            "Digital India", "Bhashini", "MeghRaj", "API Setu", "MSME"
        ]
    }
    
    print("\n--- KEYWORD DENSITY & VOCABULARY AUDIT ---")
    keyword_score = 0
    for group_name, kws in keyword_groups.items():
        found = [kw for kw in kws if re.search(r"\b" + re.escape(kw) + r"\b", text_content, re.IGNORECASE)]
        ratio = len(found) / len(kws)
        pts = round(ratio * 10)
        keyword_score += pts
        status = "[PASS]" if ratio >= 0.7 else ("[WARN]" if ratio >= 0.4 else "[FAIL]")
        print(f"  {status} {group_name}: Found {len(found)}/{len(kws)} keywords -> {pts}/10 pts")
        if len(found) < len(kws):
            missing = [kw for kw in kws if kw not in found]
            print(f"       Missing: {', '.join(missing[:4])}")

    total_score = page_score + text_score + heading_score + keyword_score
    print("\n" + "=" * 70)
    print(f"TOTAL ATS COMPLIANCE SCORE: {total_score} / 100")
    print("=" * 70)
    if total_score >= 90:
        print("RESULT: [100% ATS-PROOF] This presentation will effortlessly rank in the top percentile of automated screening.")
    elif total_score >= 75:
        print("RESULT: [STRONG PASS] Passes ATS filters, but minor keyword/heading adjustments will guarantee top rank.")
    else:
        print("RESULT: [HIGH REJECTION RISK] Presentation requires immediate text-layer and heading formatting fixes.")
    print("=" * 70)
    return total_score

if __name__ == "__main__":
    if len(sys.argv) > 1:
        target = sys.argv[1]
    else:
        target = r"d:\Aaditya\Desktop\SIH\docs\BIS_AI_SIH2026_ZenicX_Presentation.pptx"
    evaluate_ats_compliance(target)
